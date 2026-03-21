import os, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'allocate.settings')
django.setup()

from tender.models import Allocation
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

print("Fetching allocations...")

allocations = Allocation.objects.select_related(
    'job_activity__job__farmer',
    'job_activity__activity',
    'job_activity__plot',
    'mukkadam',
    'cluster',
).order_by('allocated_date')

issues = []

for alloc in allocations:
    ja = alloc.job_activity
    problems = []

    if ja.scheduled_date and alloc.allocated_date:
        if str(ja.scheduled_date) != str(alloc.allocated_date):
            problems.append(f"DATE MISMATCH: scheduled={ja.scheduled_date} but allocated={alloc.allocated_date}")

    if ja.total_area and alloc.allocated_area:
        if alloc.allocated_area > ja.total_area:
            problems.append(f"AREA OVERALLOC: total={ja.total_area}ac allocated={alloc.allocated_area}ac (over by {float(alloc.allocated_area - ja.total_area):.2f}ac)")

    if alloc.allocated_area <= Decimal('0'):
        problems.append(f"ZERO AREA: allocated_area is {alloc.allocated_area}")

    if ja.is_lost:
        problems.append("LOST ACTIVITY: allocation exists on a lost/cancelled activity")

    if not ja.scheduled_date:
        problems.append("NO SCHEDULED DATE: job activity has no scheduled_date")

    if ja.rate_per_acre and alloc.farmer_rate:
        if abs(float(alloc.farmer_rate) - float(ja.rate_per_acre)) > 0.01:
            problems.append(f"RATE MISMATCH: activity=₹{ja.rate_per_acre}/ac allocation=₹{alloc.farmer_rate}/ac")

    if problems:
        for p in problems:
            issues.append({
                'Alloc ID':         alloc.id,
                'Job ID':           ja.job.job_id,
                'Farmer':           ja.job.farmer.farmer_name,
                'Activity':         ja.activity.name,
                'Plot':             ja.plot.name if ja.plot else '—',
                'Cluster':          alloc.cluster.name if alloc.cluster else '—',
                'Mukkadam':         alloc.mukkadam.mukkadam_name,
                'Scheduled Date':   str(ja.scheduled_date) if ja.scheduled_date else 'None',
                'Allocated Date':   str(alloc.allocated_date),
                'Total Area (ac)':  float(ja.total_area),
                'Allocated Area (ac)': float(alloc.allocated_area),
                'Remaining Area (ac)': float(ja.remaining_area),
                'Activity Status':  ja.allocation_status,
                'Work Status':      alloc.work_status,
                'Farmer Rate':      float(ja.rate_per_acre or 0),
                'Alloc Farmer Rate': float(alloc.farmer_rate or 0),
                'Mukkadam Rate':    float(alloc.mukkadam_rate or 0),
                'Issue Type':       p.split(':')[0],
                'Issue Detail':     p,
            })

print(f"Found {len(issues)} issues across {len(set(i['Alloc ID'] for i in issues))} allocations")

# ── Build Excel ──────────────────────────────────────────────────────────────
wb = Workbook()

# ── Colors ───────────────────────────────────────────────────────────────────
RED_FILL    = PatternFill('solid', start_color='FFCCCC')
ORANGE_FILL = PatternFill('solid', start_color='FFE5CC')
YELLOW_FILL = PatternFill('solid', start_color='FFF9CC')
GREEN_FILL  = PatternFill('solid', start_color='CCFFCC')
HEADER_FILL = PatternFill('solid', start_color='1F3864')
ALT_FILL    = PatternFill('solid', start_color='F2F2F2')

HEADER_FONT  = Font(bold=True, color='FFFFFF', name='Arial', size=10)
BOLD_FONT    = Font(bold=True, name='Arial', size=9)
NORMAL_FONT  = Font(name='Arial', size=9)
RED_FONT     = Font(bold=True, color='CC0000', name='Arial', size=9)
ORANGE_FONT  = Font(bold=True, color='CC5500', name='Arial', size=9)

thin = Side(style='thin', color='CCCCCC')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT   = Alignment(horizontal='left',   vertical='center', wrap_text=True)

def style_header(cell):
    cell.font      = HEADER_FONT
    cell.fill      = HEADER_FILL
    cell.alignment = CENTER
    cell.border    = BORDER

def style_cell(cell, fill=None, font=None, align=None):
    cell.font      = font or NORMAL_FONT
    cell.fill      = fill or PatternFill()
    cell.alignment = align or LEFT
    cell.border    = BORDER

# ═══════════════════════════════════════════════════════════════
# SHEET 1 — Summary Dashboard
# ═══════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = 'Summary'

ws1.merge_cells('A1:D1')
ws1['A1'] = '📊 Allocation Audit Report'
ws1['A1'].font = Font(bold=True, size=16, name='Arial', color='1F3864')
ws1['A1'].alignment = CENTER

ws1['A2'] = f'Generated: {datetime.now().strftime("%d %b %Y %H:%M")}'
ws1['A2'].font = Font(italic=True, size=9, name='Arial', color='666666')

ws1.append([])

# Stats
total_allocs = allocations.count()
unique_issues = len(set(i['Alloc ID'] for i in issues))

date_issues  = [i for i in issues if 'DATE'  in i['Issue Type']]
area_issues  = [i for i in issues if 'AREA'  in i['Issue Type']]
rate_issues  = [i for i in issues if 'RATE'  in i['Issue Type']]
lost_issues  = [i for i in issues if 'LOST'  in i['Issue Type']]
zero_issues  = [i for i in issues if 'ZERO'  in i['Issue Type']]
other_issues = [i for i in issues if 'NO SC' in i['Issue Type']]

stats = [
    ['Metric',                     'Count',         'Fill'],
    ['Total Allocations Checked',  total_allocs,    None],
    ['Allocations with Issues',    unique_issues,   RED_FILL   if unique_issues > 0 else GREEN_FILL],
    ['📅 Date Mismatches',         len(set(i['Alloc ID'] for i in date_issues)),  ORANGE_FILL],
    ['📐 Area Over-Allocated',     len(set(i['Alloc ID'] for i in area_issues)),  RED_FILL],
    ['💰 Rate Mismatches',         len(set(i['Alloc ID'] for i in rate_issues)),  YELLOW_FILL],
    ['🚫 Lost Activity Allocs',    len(set(i['Alloc ID'] for i in lost_issues)),  RED_FILL],
    ['⚠️ Zero Area Allocs',        len(set(i['Alloc ID'] for i in zero_issues)),  ORANGE_FILL],
    ['❌ No Scheduled Date',       len(set(i['Alloc ID'] for i in other_issues)), YELLOW_FILL],
]

for i, row in enumerate(stats):
    r = ws1.row_dimensions[i + 4]
    r.height = 22
    a = ws1.cell(row=i+4, column=1, value=row[0])
    b = ws1.cell(row=i+4, column=2, value=row[1])
    if i == 0:
        style_header(a); style_header(b)
    else:
        fill = row[2] if row[2] else (ALT_FILL if i % 2 == 0 else PatternFill())
        style_cell(a, font=BOLD_FONT if i == 1 else NORMAL_FONT)
        style_cell(b, fill=fill, align=CENTER,
                   font=BOLD_FONT if (row[1] or 0) > 0 else NORMAL_FONT)

ws1.column_dimensions['A'].width = 35
ws1.column_dimensions['B'].width = 20

# ═══════════════════════════════════════════════════════════════
# SHEET 2 — All Issues
# ═══════════════════════════════════════════════════════════════
ws2 = wb.create_sheet('All Issues')

headers = [
    'Alloc ID', 'Job ID', 'Farmer', 'Activity', 'Plot', 'Cluster',
    'Mukkadam', 'Scheduled Date', 'Allocated Date', 'Total Area',
    'Allocated Area', 'Remaining Area', 'Activity Status', 'Work Status',
    'Farmer Rate', 'Alloc Rate', 'Mukkadam Rate', 'Issue Type', 'Issue Detail'
]

ws2.row_dimensions[1].height = 28
for col, h in enumerate(headers, 1):
    c = ws2.cell(row=1, column=col, value=h)
    style_header(c)

ISSUE_COLOR = {
    'DATE MISMATCH':  ORANGE_FILL,
    'AREA OVERALLOC': RED_FILL,
    'RATE MISMATCH':  YELLOW_FILL,
    'LOST ACTIVITY':  RED_FILL,
    'ZERO AREA':      ORANGE_FILL,
    'NO SCHEDULED':   YELLOW_FILL,
}

for row_idx, issue in enumerate(issues, 2):
    ws2.row_dimensions[row_idx].height = 20
    vals = [
        issue['Alloc ID'], issue['Job ID'], issue['Farmer'], issue['Activity'],
        issue['Plot'], issue['Cluster'], issue['Mukkadam'],
        issue['Scheduled Date'], issue['Allocated Date'],
        issue['Total Area (ac)'], issue['Allocated Area (ac)'], issue['Remaining Area (ac)'],
        issue['Activity Status'], issue['Work Status'],
        issue['Farmer Rate'], issue['Alloc Farmer Rate'], issue['Mukkadam Rate'],
        issue['Issue Type'], issue['Issue Detail'],
    ]
    row_fill = ISSUE_COLOR.get(issue['Issue Type'], ALT_FILL if row_idx % 2 == 0 else PatternFill())
    for col, val in enumerate(vals, 1):
        c = ws2.cell(row=row_idx, column=col, value=val)
        is_issue_col = col >= 18
        style_cell(c,
            fill=row_fill if is_issue_col else (ALT_FILL if row_idx % 2 == 0 else PatternFill()),
            font=RED_FONT if (is_issue_col and 'MISMATCH' in str(issue['Issue Type'])) else NORMAL_FONT,
            align=CENTER if col in [1, 8, 9, 10, 11, 12, 15, 16, 17] else LEFT
        )

col_widths = [10, 10, 28, 35, 18, 15, 22, 15, 15, 12, 14, 14, 16, 14, 12, 12, 14, 18, 55]
for i, w in enumerate(col_widths, 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

ws2.freeze_panes = 'A2'
ws2.auto_filter.ref = f'A1:{get_column_letter(len(headers))}1'

# ═══════════════════════════════════════════════════════════════
# SHEET 3 — Date Mismatches only
# ═══════════════════════════════════════════════════════════════
if date_issues:
    ws3 = wb.create_sheet('Date Mismatches')
    dh = ['Alloc ID', 'Job ID', 'Farmer', 'Activity', 'Plot', 'Cluster', 'Mukkadam',
          'Scheduled Date', 'Allocated Date', 'Days Diff', 'Total Area', 'Allocated Area', 'Work Status']
    ws3.row_dimensions[1].height = 28
    for col, h in enumerate(dh, 1):
        style_header(ws3.cell(row=1, column=col, value=h))

    seen = set()
    r = 2
    for issue in date_issues:
        if issue['Alloc ID'] in seen:
            continue
        seen.add(issue['Alloc ID'])
        ws3.row_dimensions[r].height = 20
        try:
            from datetime import date as dt
            sched = dt.fromisoformat(issue['Scheduled Date']) if issue['Scheduled Date'] != 'None' else None
            alloc = dt.fromisoformat(issue['Allocated Date'])
            diff  = (alloc - sched).days if sched else '—'
        except:
            diff = '—'

        vals = [
            issue['Alloc ID'], issue['Job ID'], issue['Farmer'], issue['Activity'],
            issue['Plot'], issue['Cluster'], issue['Mukkadam'],
            issue['Scheduled Date'], issue['Allocated Date'], diff,
            issue['Total Area (ac)'], issue['Allocated Area (ac)'], issue['Work Status'],
        ]
        fill = RED_FILL if isinstance(diff, int) and abs(diff) > 7 else ORANGE_FILL
        for col, val in enumerate(vals, 1):
            c = ws3.cell(row=r, column=col, value=val)
            style_cell(c, fill=fill if col in [8,9,10] else (ALT_FILL if r%2==0 else PatternFill()),
                       font=RED_FONT if col == 10 and isinstance(diff, int) and abs(diff) > 7 else NORMAL_FONT,
                       align=CENTER if col in [1,8,9,10,11,12] else LEFT)
        r += 1

    dw = [10,10,28,35,18,15,22,14,14,10,12,14,14]
    for i, w in enumerate(dw, 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws3.freeze_panes = 'A2'

# ═══════════════════════════════════════════════════════════════
# SHEET 4 — Area Issues
# ═══════════════════════════════════════════════════════════════
if area_issues:
    ws4 = wb.create_sheet('Area Issues')
    ah = ['Alloc ID', 'Job ID', 'Farmer', 'Activity', 'Plot', 'Cluster', 'Mukkadam',
          'Total Area', 'Allocated Area', 'Over By', 'Scheduled Date', 'Allocated Date', 'Activity Status']
    ws4.row_dimensions[1].height = 28
    for col, h in enumerate(ah, 1):
        style_header(ws4.cell(row=1, column=col, value=h))

    seen = set()
    r = 2
    for issue in area_issues:
        if issue['Alloc ID'] in seen:
            continue
        seen.add(issue['Alloc ID'])
        ws4.row_dimensions[r].height = 20
        over = round(issue['Allocated Area (ac)'] - issue['Total Area (ac)'], 2)
        vals = [
            issue['Alloc ID'], issue['Job ID'], issue['Farmer'], issue['Activity'],
            issue['Plot'], issue['Cluster'], issue['Mukkadam'],
            issue['Total Area (ac)'], issue['Allocated Area (ac)'], over,
            issue['Scheduled Date'], issue['Allocated Date'], issue['Activity Status'],
        ]
        for col, val in enumerate(vals, 1):
            c = ws4.cell(row=r, column=col, value=val)
            style_cell(c,
                fill=RED_FILL if col in [9,10] else (ALT_FILL if r%2==0 else PatternFill()),
                font=RED_FONT if col == 10 else NORMAL_FONT,
                align=CENTER if col in [1,8,9,10,11,12] else LEFT)
        r += 1

    for i, w in enumerate([10,10,28,35,18,15,22,12,14,10,14,14,16], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.freeze_panes = 'A2'

# With:
import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), f'allocation_audit_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx')
wb.save(out)
print(f'\n✅ Saved to: {out}')