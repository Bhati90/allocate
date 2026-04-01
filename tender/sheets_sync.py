"""
Django Management Command: fixsales
====================================
Place at: tender/management/commands/fixsales.py

Rules:
1. Allocated → NEVER touch
2. 0-acre → SKIP
3. If Pruning scheduled_date > Pruning sales_date (moved forward):
   → Cascade: all unallocated activities in that plot get scheduled_date = pruning_scheduled + gap
4. If Pruning scheduled_date <= Pruning sales_date:
   → Old logic: only fix if unallocated + scheduled is more than 5 days before sales
5. No pruning in plot → no change

Usage:
    python manage.py fixsales
    python manage.py fixsales --apply
    python manage.py fixsales --cluster=Nashik
"""

from datetime import timedelta
from collections import Counter

from django.core.management.base import BaseCommand
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from tender.models import (
    Plot, JobActivity,
    ClusterActivityScheduleRule, ActivityScheduleRule,
)

HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
BODY_FONT = Font(name='Calibri', size=11)
CASCADE_FILL = PatternFill('solid', fgColor='DAEEF3')
UPDATE_FILL = PatternFill('solid', fgColor='FFF2CC')
SKIP_FILL = PatternFill('solid', fgColor='E2EFDA')
CENTER = Alignment(horizontal='center', vertical='center')
LEFT = Alignment(horizontal='left', vertical='center', indent=1)


class Command(BaseCommand):
    help = "Fix scheduled_date based on Pruning base + gap days"

    def add_arguments(self, parser):
        parser.add_argument('--apply', action='store_true', default=False)
        parser.add_argument('--cluster', type=str, default=None)
        parser.add_argument('--plot-id', type=int, default=None)
        parser.add_argument('--output', type=str, default='schedule_fix_report.xlsx')

    def handle(self, *args, **options):
        apply_changes = options['apply']
        out_path = options['output']

        if apply_changes:
            self.stdout.write(self.style.WARNING("⚠️  APPLY MODE"))
        else:
            self.stdout.write(self.style.NOTICE("🔍 DRY RUN"))

        plots_qs = Plot.objects.all().prefetch_related(
            'activities__activity', 'activities__allocations', 'clusters',
        ).select_related('farmer')

        if options['plot_id']:
            plots_qs = plots_qs.filter(id=options['plot_id'])
        if options['cluster']:
            plots_qs = plots_qs.filter(clusters__name=options['cluster']).distinct()

        results = []
        update_count = 0

        for plot in plots_qs.iterator(chunk_size=100):
            job_activities = list(
                JobActivity.objects
                .filter(plot=plot)
                .exclude(total_area=0)
                .select_related('activity', 'job')
                .order_by('scheduled_date')
            )

            pruning_ja = None
            for ja in job_activities:
                if self._is_pruning(ja.activity):
                    pruning_ja = ja
                    break

            base_date = pruning_ja.scheduled_date if pruning_ja and pruning_ja.scheduled_date else None

            pruning_moved_forward = False
            if pruning_ja and pruning_ja.scheduled_date and pruning_ja.sales_date:
                pruning_moved_forward = pruning_ja.scheduled_date > pruning_ja.sales_date

            for ja in job_activities:
                row = self._evaluate(ja, plot, base_date, pruning_moved_forward)
                results.append(row)

                if apply_changes and row['action'] in ('CASCADE from pruning move', 'UPDATE scheduled_date'):
                    new_date = row['_new_date_raw']
                    if new_date:
                        ja.scheduled_date = new_date
                        ja.save(update_fields=['scheduled_date', 'updated_at'])
                        update_count += 1

        self._write_xlsx(results, out_path)

        actions = Counter(r['action'] for r in results)
        self.stdout.write(f"\nTotal: {len(results)}")
        for act, cnt in actions.most_common():
            self.stdout.write(f"  {act}: {cnt}")
        self.stdout.write(f"\nReport: {out_path}")

        if apply_changes:
            self.stdout.write(self.style.SUCCESS(f"✅ Updated {update_count} scheduled_dates."))
        else:
            needs = sum(1 for r in results if r['action'] in ('CASCADE from pruning move', 'UPDATE scheduled_date'))
            self.stdout.write(self.style.WARNING(f"⏳ {needs} need updating. Re-run with --apply"))

    def _write_xlsx(self, results, path):
        wb = Workbook()
        ws = wb.active
        ws.title = "Schedule Fix Report"

        headers = [
            'Farmer', 'Plot', 'Acres', 'Activity', 'Current Scheduled',
            'Sales Date', 'Diff (days)', 'Allocated?', 'Pruning Moved?',
            'Gap Days', 'Gap Source', 'New Scheduled', 'Action',
        ]

        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.fill = HDR_FILL
            cell.font = HDR_FONT
            cell.alignment = CENTER

        for i, r in enumerate(results, 2):
            vals = [
                r['farmer'], r['plot_name'], r['acres'], r['activity'],
                r['current_scheduled_date'], r['sales_date'], r['diff_days'],
                'Yes' if r['is_allocated'] else 'No',
                'Yes' if r['pruning_moved'] else 'No',
                r['gap_days_used'], r['gap_source'],
                r['new_scheduled_date'], r['action'],
            ]
            action = r['action']
            is_cascade = action == 'CASCADE from pruning move'
            is_update = action == 'UPDATE scheduled_date'
            is_skip = 'SKIP' in action

            for c, v in enumerate(vals, 1):
                cell = ws.cell(row=i, column=c, value=v)
                cell.font = BODY_FONT
                if c in (3, 7, 10):
                    cell.alignment = CENTER
                    if c == 3:
                        cell.number_format = '0.00'
                elif c in (5, 6, 8, 9, 12):
                    cell.alignment = CENTER
                else:
                    cell.alignment = LEFT

                if is_cascade:
                    cell.fill = CASCADE_FILL
                elif is_update:
                    cell.fill = UPDATE_FILL
                elif is_skip:
                    cell.fill = SKIP_FILL

        widths = [20, 18, 8, 25, 16, 16, 10, 10, 12, 9, 20, 16, 30]
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(results)+1}"

        wb.save(path)

    @staticmethod
    def _is_pruning(activity):
        name_lower = activity.name.lower()
        return 'pruning' in name_lower or 'छाटणी' in activity.name

    @staticmethod
    def _get_gap_days(plot, activity):
        for cluster in plot.clusters.all():
            rule = ClusterActivityScheduleRule.objects.filter(
                cluster=cluster, activity=activity,
            ).first()
            if rule:
                return rule.gap_days, f"Cluster: {cluster.name}"

        global_rule = ActivityScheduleRule.objects.filter(activity=activity).first()
        if global_rule:
            return global_rule.gap_days, "Global Rule"

        default = activity.default_gap_days or 3
        return default, "Catalog Default" if activity.default_gap_days else "Fallback (3)"

    def _evaluate(self, ja, plot, base_date, pruning_moved_forward):
        current_sched = ja.scheduled_date
        current_sales = ja.sales_date
        is_alloc = ja.allocation_status in ('fully_allocated', 'completed', 'in_progress')
        is_pruning = self._is_pruning(ja.activity)

        gap_days_used = None
        gap_source = ""
        new_sched = current_sched
        action = "NO CHANGE"

        if is_pruning:
            expected = current_sched
            gap_days_used = 0
            gap_source = "BASE (Pruning)"
        elif base_date is None:
            expected = current_sched
            gap_source = "No Pruning in plot"
        else:
            gap_days_used, gap_source = self._get_gap_days(plot, ja.activity)
            expected = base_date + timedelta(days=gap_days_used)

        if is_alloc:
            action = "SKIP (allocated)"
        elif is_pruning:
            action = "BASE (Pruning)"
        elif base_date is None:
            action = "OK (no pruning in plot)"
        elif pruning_moved_forward:
            if current_sched != expected:
                action = "CASCADE from pruning move"
                new_sched = expected
            else:
                action = "OK (already correct after cascade)"
        elif current_sched and current_sales:
            diff = (current_sched - current_sales).days
            if diff < -5:
                action = "UPDATE scheduled_date"
                new_sched = expected
            else:
                action = f"OK (diff={diff}d, within tolerance)"
        elif current_sched and not current_sales:
            action = "OK (no sales_date to compare)"
        else:
            action = "OK (no scheduled_date)"

        diff_days = (current_sched - current_sales).days if current_sched and current_sales else None
        farmer_name = getattr(plot.farmer, 'farmer_name', '?')
        plot_acres = float(plot.area_acres) if plot.area_acres else 0

        return {
            'farmer': farmer_name,
            'plot_name': plot.name,
            'acres': plot_acres,
            'activity': ja.activity.name,
            'current_scheduled_date': current_sched,
            'sales_date': current_sales,
            'diff_days': diff_days,
            'is_allocated': is_alloc,
            'pruning_moved': pruning_moved_forward,
            'gap_days_used': gap_days_used,
            'gap_source': gap_source,
            'new_scheduled_date': new_sched,
            'action': action,
            '_new_date_raw': new_sched,
        }