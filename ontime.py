import os, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'allocate.settings')
django.setup()

from datetime import date, datetime, timedelta
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from tender.models import JobActivity

print("Fetching data...")

qs = JobActivity.objects.filter(
    is_lost=False,
    total_area__gt=0,
    sales_date__isnull=False,
    allocations__work_status='completed',
).distinct().select_related(
    'activity', 'job__farmer', 'plot'
).prefetch_related('allocations', 'plot__clusters', 'job__clusters')

# ── GRANULAR BUCKETS ─────────────────────────────────────────────────
on_time     = []
late_1d     = []
late_2d     = []
late_3d     = []
late_4_6d   = []
late_7_10d  = []
late_11_15d = []
late_16_30d = []
late_30p    = []

# ── Deduplicate by (job_id, plot_id, activity_name) ──────────────────
seen = {}  # dedup_key → item

for ja in qs:
    completed_alloc = ja.allocations.filter(
        work_status='completed'
    ).order_by('allocated_date').first()  # earliest completed
    if not completed_alloc:
        continue

    if completed_alloc.actual_end_time:
        completed_date = completed_alloc.actual_end_time.date()
    elif completed_alloc.report_submitted_at:
        completed_date = completed_alloc.report_submitted_at.date()
    elif completed_alloc.allocated_date:
        completed_date = completed_alloc.allocated_date
    else:
        continue

    sales_date = ja.sales_date
    diff       = (completed_date - sales_date).days

    cluster = (
        ja.plot.clusters.first().name if ja.plot and ja.plot.clusters.exists() else
        ja.job.clusters.first().name  if ja.job.clusters.exists() else '—'
    )

    item = {
        'ja_id':          ja.id,
        'job_id':         ja.job.job_id,
        'farmer':         ja.job.farmer.farmer_name,
        'activity':       ja.activity.name,
        'plot':           ja.plot.name if ja.plot else '—',
        'cluster':        cluster,
        'scheduled_date': str(ja.scheduled_date) if ja.scheduled_date else '—',
        'sales_date':     str(sales_date),
        'completed_date': str(completed_date),
        'diff_days':      diff,
        'total_area':     float(ja.total_area),
    }

    # ── One entry per job + plot + activity — keep best (lowest diff) ──
    dedup_key = (str(ja.job.job_id), str(ja.plot_id or ''), ja.activity.name)
    if dedup_key not in seen or diff < seen[dedup_key]['diff_days']:
        seen[dedup_key] = item

# ── Bucket deduplicated items ─────────────────────────────────────────
on_time     = []
late_1d     = []
late_2d     = []
late_3d     = []
late_4_6d   = []
late_7_10d  = []
late_11_15d = []
late_16_30d = []
late_30p    = []

for item in seen.values():
    diff = item['diff_days']
    if diff <= 0:       on_time.append(item)
    elif diff == 1:     late_1d.append(item)
    elif diff == 2:     late_2d.append(item)
    elif diff == 3:     late_3d.append(item)
    elif diff <= 6:     late_4_6d.append(item)
    elif diff <= 10:    late_7_10d.append(item)
    elif diff <= 15:    late_11_15d.append(item)
    elif diff <= 30:    late_16_30d.append(item)
    else:               late_30p.append(item)
late_1_3d = late_1d + late_2d + late_3d
all_buckets = [on_time, late_1d, late_2d, late_3d, late_4_6d,
               late_7_10d, late_11_15d, late_16_30d, late_30p]
all_items = [i for b in all_buckets for i in b]
counted = len(all_items)
total_area = sum(i['total_area'] for i in all_items)
print(f"Analysed {counted} activities | Total area: {round(total_area, 2)} ac")

def pct(n, d=None):
    d = d or counted
    return round(n / d * 100, 1) if d > 0 else 0


# ── Styles ───────────────────────────────────────────────────────────
HEADER_FILL   = PatternFill('solid', start_color='1F3864')
GREEN_FILL    = PatternFill('solid', start_color='C6EFCE')
LIME_FILL     = PatternFill('solid', start_color='E2EFDA')
YELLOW_FILL   = PatternFill('solid', start_color='FFEB9C')
GOLD_FILL     = PatternFill('solid', start_color='FFD966')
ORANGE_FILL   = PatternFill('solid', start_color='FFCC99')
DEEP_ORANGE   = PatternFill('solid', start_color='F4B084')
RED_FILL      = PatternFill('solid', start_color='FFC7CE')
DARK_RED_FILL = PatternFill('solid', start_color='FF0000')
ALT_FILL      = PatternFill('solid', start_color='F5F5F5')

HEADER_FONT  = Font(bold=True, color='FFFFFF', name='Arial', size=10)
BOLD_FONT    = Font(bold=True, name='Arial', size=10)
NORMAL_FONT  = Font(name='Arial', size=9)
TITLE_FONT   = Font(bold=True, size=12, name='Arial', color='1F3864')

thin   = Side(style='thin', color='CCCCCC')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal='center', vertical='center')
LEFT   = Alignment(horizontal='left', vertical='center')
WRAP   = Alignment(horizontal='center', vertical='center', wrap_text=True)

BUCKET_STYLES = [
    ('✅ On Time (≤ 0d)',  on_time,     '375623', 'C6EFCE', GREEN_FILL),
    ('🟢 Late 1d',         late_1d,     '375623', 'E2EFDA', LIME_FILL),
    ('🟡 Late 2d',         late_2d,     '7D6608', 'FFEB9C', YELLOW_FILL),
    ('🟡 Late 3d',         late_3d,     '7D6608', 'FFD966', GOLD_FILL),
    ('🟠 Late 4–6d',       late_4_6d,   '833C00', 'FFCC99', ORANGE_FILL),
    ('🔶 Late 7–10d',      late_7_10d,  '833C00', 'F4B084', DEEP_ORANGE),
    ('🔴 Late 11–15d',     late_11_15d, '9C0006', 'FFC7CE', RED_FILL),
    ('🔴 Late 16–30d',     late_16_30d, '9C0006', 'FFC7CE', RED_FILL),
    ('💀 Late 30d+',       late_30p,    'FFFFFF', 'FF0000', DARK_RED_FILL),
]

ROLLUP_STYLES = [
    ('✅ On Time (≤ 0d)',  on_time,     '375623', 'C6EFCE', GREEN_FILL),
    ('🟢 Late 1–3d',       late_1_3d,   '375623', 'E2EFDA', LIME_FILL),
    ('🟠 Late 4–6d',       late_4_6d,   '833C00', 'FFCC99', ORANGE_FILL),
    ('🔶 Late 7–10d',      late_7_10d,  '833C00', 'F4B084', DEEP_ORANGE),
    ('🔴 Late 11–15d',     late_11_15d, '9C0006', 'FFC7CE', RED_FILL),
    ('🔴 Late 16–30d',     late_16_30d, '9C0006', 'FFC7CE', RED_FILL),
    ('💀 Late 30d+',       late_30p,    'FFFFFF', 'FF0000', DARK_RED_FILL),
]


def style_header(cell):
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = CENTER
    cell.border = BORDER

def style_cell(cell, fill=None, font=None, align=None):
    cell.font = font or NORMAL_FONT
    cell.fill = fill or PatternFill()
    cell.alignment = align or LEFT
    cell.border = BORDER

def get_diff_fill(d):
    if d <= 0:  return GREEN_FILL
    if d <= 3:  return LIME_FILL
    if d <= 6:  return ORANGE_FILL
    if d <= 10: return DEEP_ORANGE
    if d <= 15: return RED_FILL
    if d <= 30: return RED_FILL
    return DARK_RED_FILL

def pct_fill(v, good=70, mid=40):
    if v >= good: return GREEN_FILL
    if v >= mid:  return YELLOW_FILL
    return RED_FILL

def write_bucket_table(ws, row, bucket_list, title):
    ws.cell(row=row, column=1, value=title).font = TITLE_FONT
    row += 1
    headers = ['Bucket', 'Count', '% Share', 'Avg Days Late',
               'Total Area (ac)', 'Area-Wtd Avg Days']
    ws.row_dimensions[row].height = 22
    for col, h in enumerate(headers, 1):
        style_header(ws.cell(row=row, column=col, value=h))
    row += 1

    for label, items, text_color, bg_color, fill in bucket_list:
        ws.row_dimensions[row].height = 20
        n = len(items)
        avg_diff = round(sum(i['diff_days'] for i in items) / n, 1) if n else 0
        tot_ac = round(sum(i['total_area'] for i in items), 2)
        area_wtd = round(
            sum(i['diff_days'] * i['total_area'] for i in items) / tot_ac, 1
        ) if tot_ac > 0 else 0
        vals = [label, n, f'{pct(n)}%', avg_diff, tot_ac, area_wtd]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=row, column=col, value=val)
            txt = Font(bold=True,
                       color='FFFFFF' if label.startswith('💀') else text_color,
                       name='Arial', size=10)
            style_cell(c, fill=fill, font=txt,
                       align=CENTER if col > 1 else LEFT)
        row += 1

    ws.row_dimensions[row].height = 22
    area_wtd_all = round(
        sum(i['diff_days'] * i['total_area'] for i in all_items) / total_area, 1
    ) if total_area > 0 else 0
    tot_vals = ['TOTAL', counted, '100%',
                round(sum(i['diff_days'] for i in all_items) / counted, 1) if counted else 0,
                round(total_area, 2), area_wtd_all]
    for col, val in enumerate(tot_vals, 1):
        c = ws.cell(row=row, column=col, value=val)
        c.font = BOLD_FONT
        c.fill = PatternFill('solid', start_color='D9D9D9')
        c.alignment = CENTER if col > 1 else LEFT
        c.border = BORDER
    return row + 2


wb = Workbook()

# ═══════════════════════════════════════════════════════════════════════
# SHEET 1 — Summary Dashboard
# ═══════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = 'Summary'

ws1.merge_cells('A1:G1')
ws1['A1'] = '📊 Sales Date Performance Report'
ws1['A1'].font = Font(bold=True, size=16, name='Arial', color='1F3864')
ws1['A1'].alignment = CENTER

ot_count = len(on_time)
within_3 = ot_count + len(late_1_3d)
ot_pct_val = pct(ot_count)
w3_pct_val = pct(within_3)
ot_area = sum(i['total_area'] for i in on_time)
w3_area = sum(i['total_area'] for i in on_time + late_1_3d)
ot_area_pct = pct(ot_area, total_area)
w3_area_pct = pct(w3_area, total_area)

ws1['A2'] = (f'Generated: {datetime.now().strftime("%d %b %Y %H:%M")} | '
             f'Activities: {counted} | Area: {round(total_area,1)} ac')
ws1['A2'].font = Font(italic=True, size=9, name='Arial', color='666666')

kpi_row = 3
kpi_data = [
    ('On Time %', f'{ot_pct_val}%', pct_fill(ot_pct_val)),
    ('≤ 3d %', f'{w3_pct_val}%', pct_fill(w3_pct_val, 85, 60)),
    ('On Time % (area)', f'{ot_area_pct}%', pct_fill(ot_area_pct)),
    ('≤ 3d % (area)', f'{w3_area_pct}%', pct_fill(w3_area_pct, 85, 60)),
]
for col, (lbl, val, fill) in enumerate(kpi_data, 1):
    c = ws1.cell(row=kpi_row, column=col * 2 - 1, value=lbl)
    c.font = Font(bold=True, name='Arial', size=9, color='666666')
    c.alignment = CENTER
    c.border = BORDER
    c = ws1.cell(row=kpi_row, column=col * 2, value=val)
    c.font = Font(bold=True, name='Arial', size=14)
    c.fill = fill
    c.alignment = CENTER
    c.border = BORDER

row = 5
row = write_bucket_table(ws1, row, BUCKET_STYLES, 'DAY-LEVEL BREAKDOWN')
row = write_bucket_table(ws1, row, ROLLUP_STYLES, 'RANGE ROLL-UP')

# Per-cluster breakdown
ws1.cell(row=row, column=1, value='PER-CLUSTER BREAKDOWN').font = TITLE_FONT
row += 1

cl_headers = ['Cluster', 'Total', 'On Time', '1-3d', '4-6d', '7-10d',
              '11-15d', '16-30d', '30d+', 'On Time %', '≤3d %',
              'Area (ac)', 'OT % (area)']
ws1.row_dimensions[row].height = 28
for col, h in enumerate(cl_headers, 1):
    c = ws1.cell(row=row, column=col, value=h)
    style_header(c)
    c.alignment = WRAP
row += 1

bucket_key_map = {
    '✅ On Time (≤ 0d)': 'ot',
    '🟢 Late 1d': 'l1_3', '🟡 Late 2d': 'l1_3', '🟡 Late 3d': 'l1_3',
    '🟠 Late 4–6d': 'l4_6', '🔶 Late 7–10d': 'l7_10',
    '🔴 Late 11–15d': 'l11_15', '🔴 Late 16–30d': 'l16_30',
    '💀 Late 30d+': 'l30p',
}
cs = defaultdict(lambda: {k: 0 for k in
    ['ot', 'l1_3', 'l4_6', 'l7_10', 'l11_15', 'l16_30', 'l30p']}
    | {'area': 0.0, 'ot_area': 0.0, 'w3_area': 0.0})
for label, items, *_ in BUCKET_STYLES:
    key = bucket_key_map[label]
    for it in items:
        cs[it['cluster']][key] += 1
        cs[it['cluster']]['area'] += it['total_area']
        if key == 'ot':
            cs[it['cluster']]['ot_area'] += it['total_area']
        if key in ('ot', 'l1_3'):
            cs[it['cluster']]['w3_area'] += it['total_area']

for ri, (cl, s) in enumerate(sorted(cs.items(),
        key=lambda x: -(sum(v for k, v in x[1].items()
                            if k not in ('area', 'ot_area', 'w3_area'))))):
    tot = sum(v for k, v in s.items() if k not in ('area', 'ot_area', 'w3_area'))
    ot_p = pct(s['ot'], tot)
    w3_p = pct(s['ot'] + s['l1_3'], tot)
    ot_a_p = pct(s['ot_area'], s['area']) if s['area'] else 0
    fill = ALT_FILL if ri % 2 == 0 else PatternFill()
    vals = [cl, tot, s['ot'], s['l1_3'], s['l4_6'], s['l7_10'],
            s['l11_15'], s['l16_30'], s['l30p'],
            f'{ot_p}%', f'{w3_p}%', round(s['area'], 1), f'{ot_a_p}%']
    ws1.row_dimensions[row].height = 18
    for col, val in enumerate(vals, 1):
        c = ws1.cell(row=row, column=col, value=val)
        cf = (pct_fill(ot_p) if col == 10
              else pct_fill(w3_p, 85, 60) if col == 11
              else pct_fill(ot_a_p) if col == 13
              else fill)
        style_cell(c, fill=cf,
                   font=BOLD_FONT if col in [1, 2, 10, 11, 13] else NORMAL_FONT,
                   align=CENTER if col > 1 else LEFT)
    row += 1

for i, w in enumerate([26, 8, 9, 8, 8, 8, 8, 8, 8, 10, 10, 11, 11], 1):
    ws1.column_dimensions[get_column_letter(i)].width = w
ws1.freeze_panes = 'A5'


# ═══════════════════════════════════════════════════════════════════════
# SHEET 2 — Weekly Trend
# ═══════════════════════════════════════════════════════════════════════
ws_trend = wb.create_sheet('Weekly Trend')
ws_trend.merge_cells('A1:L1')
ws_trend['A1'] = '📈 Week-over-Week Performance (cohort by Sales Date week)'
ws_trend['A1'].font = Font(bold=True, size=14, name='Arial', color='1F3864')
ws_trend['A1'].alignment = CENTER

week_stats = defaultdict(lambda: {
    'total': 0, 'ot': 0, 'w3': 0,
    'area': 0.0, 'ot_area': 0.0, 'w3_area': 0.0, 'diff_sum': 0
})
for it in all_items:
    sd = (datetime.strptime(it['sales_date'], '%Y-%m-%d').date()
          if isinstance(it['sales_date'], str) else it['sales_date'])
    wk_start = sd - timedelta(days=sd.weekday())
    week_stats[wk_start]['total'] += 1
    week_stats[wk_start]['area'] += it['total_area']
    week_stats[wk_start]['diff_sum'] += it['diff_days']
    if it['diff_days'] <= 0:
        week_stats[wk_start]['ot'] += 1
        week_stats[wk_start]['ot_area'] += it['total_area']
    if it['diff_days'] <= 3:
        week_stats[wk_start]['w3'] += 1
        week_stats[wk_start]['w3_area'] += it['total_area']

trend_headers = ['Week Starting', 'Activities', 'On Time', 'On Time %',
                 '≤ 3d', '≤ 3d %', 'Avg Days Late',
                 'Area (ac)', 'OT % (area)', '≤ 3d % (area)',
                 'WoW Δ OT%', 'WoW Δ ≤3d%']
ws_trend.row_dimensions[2].height = 24
for col, h in enumerate(trend_headers, 1):
    c = ws_trend.cell(row=2, column=col, value=h)
    style_header(c)
    c.alignment = WRAP

prev_ot = prev_w3 = None
for ri, (wk, s) in enumerate(sorted(week_stats.items())):
    r = ri + 3
    ws_trend.row_dimensions[r].height = 18
    fill = ALT_FILL if ri % 2 == 0 else PatternFill()
    ot_p = pct(s['ot'], s['total'])
    w3_p = pct(s['w3'], s['total'])
    avg_d = round(s['diff_sum'] / s['total'], 1) if s['total'] else 0
    ot_a = pct(s['ot_area'], s['area']) if s['area'] else 0
    w3_a = pct(s['w3_area'], s['area']) if s['area'] else 0
    wow_ot = f'{round(ot_p - prev_ot, 1):+}pp' if prev_ot is not None else '—'
    wow_w3 = f'{round(w3_p - prev_w3, 1):+}pp' if prev_w3 is not None else '—'
    prev_ot, prev_w3 = ot_p, w3_p

    vals = [wk.strftime('%d %b %Y'), s['total'], s['ot'], f'{ot_p}%',
            s['w3'], f'{w3_p}%', avg_d,
            round(s['area'], 1), f'{ot_a}%', f'{w3_a}%', wow_ot, wow_w3]
    for col, val in enumerate(vals, 1):
        c = ws_trend.cell(row=r, column=col, value=val)
        cf = (pct_fill(ot_p) if col == 4
              else pct_fill(w3_p, 85, 60) if col == 6
              else pct_fill(ot_a) if col == 9
              else pct_fill(w3_a, 85, 60) if col == 10
              else fill)
        style_cell(c, fill=cf,
                   font=BOLD_FONT if col in [1, 4, 6, 9, 10] else NORMAL_FONT,
                   align=CENTER if col > 1 else LEFT)

for i, w in enumerate([14, 10, 9, 10, 8, 10, 12, 11, 11, 12, 11, 11], 1):
    ws_trend.column_dimensions[get_column_letter(i)].width = w
ws_trend.freeze_panes = 'A3'


# ═══════════════════════════════════════════════════════════════════════
# SHEET 3 — Cluster × Activity Heat Map
# ═══════════════════════════════════════════════════════════════════════
ws_heat = wb.create_sheet('Cluster × Activity')
ws_heat.merge_cells('A1:E1')
ws_heat['A1'] = '🔥 Cluster × Activity — On-Time % Heat Map'
ws_heat['A1'].font = Font(bold=True, size=14, name='Arial', color='1F3864')
ws_heat['A1'].alignment = CENTER

ca = defaultdict(lambda: defaultdict(
    lambda: {'total': 0, 'ot': 0, 'w3': 0, 'area': 0.0}))
for it in all_items:
    ca[it['cluster']][it['activity']]['total'] += 1
    ca[it['cluster']][it['activity']]['area'] += it['total_area']
    if it['diff_days'] <= 0:
        ca[it['cluster']][it['activity']]['ot'] += 1
    if it['diff_days'] <= 3:
        ca[it['cluster']][it['activity']]['w3'] += 1

activities = sorted({it['activity'] for it in all_items})
clusters_sorted = sorted(ca.keys())

heat_headers = ['Cluster'] + activities + ['Overall OT%']
ws_heat.row_dimensions[2].height = 24
for col, h in enumerate(heat_headers, 1):
    c = ws_heat.cell(row=2, column=col, value=h)
    style_header(c)
    c.alignment = WRAP

for ri, cl in enumerate(clusters_sorted):
    r = ri + 3
    ws_heat.row_dimensions[r].height = 18
    fill = ALT_FILL if ri % 2 == 0 else PatternFill()
    style_cell(ws_heat.cell(row=r, column=1, value=cl), fill=fill, font=BOLD_FONT)

    cl_total = cl_ot = 0
    for ci, act in enumerate(activities, 2):
        s = ca[cl][act]
        cl_total += s['total']
        cl_ot += s['ot']
        if s['total'] == 0:
            c = ws_heat.cell(row=r, column=ci, value='—')
            style_cell(c, fill=fill, align=CENTER)
        else:
            v = pct(s['ot'], s['total'])
            c = ws_heat.cell(row=r, column=ci, value=f'{v}% ({s["total"]})')
            style_cell(c, fill=pct_fill(v), font=BOLD_FONT, align=CENTER)

    ov = pct(cl_ot, cl_total) if cl_total else 0
    c = ws_heat.cell(row=r, column=len(heat_headers), value=f'{ov}%')
    style_cell(c, fill=pct_fill(ov), font=BOLD_FONT, align=CENTER)

ws_heat.column_dimensions['A'].width = 26
for i in range(2, len(heat_headers) + 1):
    ws_heat.column_dimensions[get_column_letter(i)].width = 18
ws_heat.freeze_panes = 'B3'


# ═══════════════════════════════════════════════════════════════════════
# SHEET 4 — Farmer Performance
# ═══════════════════════════════════════════════════════════════════════
ws_farm = wb.create_sheet('Farmer Performance')
ws_farm.merge_cells('A1:J1')
ws_farm['A1'] = '👨‍🌾 Farmer-Level Performance (sorted worst → best)'
ws_farm['A1'].font = Font(bold=True, size=14, name='Arial', color='1F3864')
ws_farm['A1'].alignment = CENTER

farmer_stats = defaultdict(lambda: {
    'total': 0, 'ot': 0, 'w3': 0,
    'area': 0.0, 'ot_area': 0.0, 'diff_sum': 0, 'clusters': set()
})
for it in all_items:
    f = it['farmer']
    farmer_stats[f]['total'] += 1
    farmer_stats[f]['area'] += it['total_area']
    farmer_stats[f]['diff_sum'] += it['diff_days']
    farmer_stats[f]['clusters'].add(it['cluster'])
    if it['diff_days'] <= 0:
        farmer_stats[f]['ot'] += 1
        farmer_stats[f]['ot_area'] += it['total_area']
    if it['diff_days'] <= 3:
        farmer_stats[f]['w3'] += 1

farm_headers = ['Farmer', 'Cluster(s)', 'Activities', 'On Time', 'On Time %',
                '≤ 3d', '≤ 3d %', 'Avg Days Late', 'Area (ac)', 'Late Area Risk']
ws_farm.row_dimensions[2].height = 22
for col, h in enumerate(farm_headers, 1):
    style_header(ws_farm.cell(row=2, column=col, value=h))

farmers_sorted = sorted(farmer_stats.items(),
                         key=lambda x: pct(x[1]['ot'], x[1]['total']))

for ri, (fname, s) in enumerate(farmers_sorted):
    r = ri + 3
    ws_farm.row_dimensions[r].height = 18
    fill = ALT_FILL if ri % 2 == 0 else PatternFill()
    ot_p = pct(s['ot'], s['total'])
    w3_p = pct(s['w3'], s['total'])
    avg_d = round(s['diff_sum'] / s['total'], 1) if s['total'] else 0
    late_area = round(s['area'] - s['ot_area'], 2)
    vals = [fname, ', '.join(sorted(s['clusters'])), s['total'],
            s['ot'], f'{ot_p}%', s['w3'], f'{w3_p}%',
            avg_d, round(s['area'], 2), late_area]
    for col, val in enumerate(vals, 1):
        c = ws_farm.cell(row=r, column=col, value=val)
        cf = (pct_fill(ot_p) if col == 5
              else pct_fill(w3_p, 85, 60) if col == 7
              else fill)
        style_cell(c, fill=cf,
                   font=BOLD_FONT if col in [1, 5, 7] else NORMAL_FONT,
                   align=CENTER if col > 2 else LEFT)

for i, w in enumerate([30, 28, 10, 9, 10, 8, 10, 13, 11, 12], 1):
    ws_farm.column_dimensions[get_column_letter(i)].width = w
ws_farm.freeze_panes = 'A3'
ws_farm.auto_filter.ref = f'A2:{get_column_letter(len(farm_headers))}2'


# ═══════════════════════════════════════════════════════════════════════
# SHEETS 5–13 — Detail per granular bucket
# ═══════════════════════════════════════════════════════════════════════
detail_headers = [
    'JA ID', 'Job ID', 'Farmer', 'Activity', 'Plot', 'Cluster',
    'Scheduled Date', 'Sales Date', 'Completed Date', 'Days Diff', 'Area (ac)'
]
col_widths_detail = [8, 10, 28, 35, 18, 20, 14, 14, 14, 10, 10]

for label, items, text_color, bg_color, fill in BUCKET_STYLES:
    sheet_name = label.split(' ', 1)[1][:28]
    ws = wb.create_sheet(sheet_name)

    ws.merge_cells(f'A1:{get_column_letter(len(detail_headers))}1')
    ws['A1'] = f'{label} — {len(items)} jobs ({pct(len(items))}%)'
    ws['A1'].font = Font(bold=True, size=13, name='Arial',
                         color='FFFFFF' if text_color == 'FFFFFF' else text_color)
    ws['A1'].fill = PatternFill('solid', start_color=bg_color)
    ws['A1'].alignment = CENTER

    ws.row_dimensions[2].height = 22
    for col, h in enumerate(detail_headers, 1):
        style_header(ws.cell(row=2, column=col, value=h))

    for ri, item in enumerate(
            sorted(items, key=lambda x: x['diff_days'], reverse=True), 3):
        ws.row_dimensions[ri].height = 18
        row_fill = ALT_FILL if ri % 2 == 0 else PatternFill()
        vals = [
            item['ja_id'], item['job_id'], item['farmer'], item['activity'],
            item['plot'], item['cluster'], item['scheduled_date'],
            item['sales_date'], item['completed_date'],
            item['diff_days'], item['total_area']
        ]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=ri, column=col, value=val)
            is_diff = col == 10
            style_cell(c,
                fill=get_diff_fill(item['diff_days']) if is_diff else row_fill,
                font=Font(bold=True, name='Arial', size=9,
                          color='FFFFFF' if item['diff_days'] > 30
                          and is_diff else '000000')
                if is_diff else NORMAL_FONT,
                align=CENTER if col in [1, 2, 7, 8, 9, 10, 11] else LEFT)

    for i, w in enumerate(col_widths_detail, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A3'
    ws.auto_filter.ref = f'A2:{get_column_letter(len(detail_headers))}2'


# ═══════════════════════════════════════════════════════════════════════
# LAST SHEET — All Data
# ═══════════════════════════════════════════════════════════════════════
ws_all = wb.create_sheet('All Data')
ws_all.row_dimensions[1].height = 22
for col, h in enumerate(detail_headers + ['Bucket'], 1):
    style_header(ws_all.cell(row=1, column=col, value=h))

ri = 2
for label, items, *_ in BUCKET_STYLES:
    bucket_name = label.split(' ', 1)[1]
    for item in sorted(items, key=lambda x: x['diff_days'], reverse=True):
        ws_all.row_dimensions[ri].height = 18
        fill = ALT_FILL if ri % 2 == 0 else PatternFill()
        vals = [
            item['ja_id'], item['job_id'], item['farmer'], item['activity'],
            item['plot'], item['cluster'], item['scheduled_date'],
            item['sales_date'], item['completed_date'],
            item['diff_days'], item['total_area'], bucket_name
        ]
        for col, val in enumerate(vals, 1):
            style_cell(ws_all.cell(row=ri, column=col, value=val),
                       fill=fill,
                       align=CENTER if col in [1, 2, 7, 8, 9, 10, 11] else LEFT)
        ri += 1

for i, w in enumerate(col_widths_detail + [18], 1):
    ws_all.column_dimensions[get_column_letter(i)].width = w
ws_all.freeze_panes = 'A2'
ws_all.auto_filter.ref = f'A1:{get_column_letter(len(detail_headers) + 1)}1'


# ── Save ─────────────────────────────────────────────────────────────
out = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    f'sales_performance_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
)
wb.save(out)
print(f'\n✅ Saved: {out}')
print(f'\n{"=" * 50}')
print(f'KPIs:  On Time = {ot_pct_val}% | ≤ 3d = {w3_pct_val}%')
print(f'       On Time (area) = {ot_area_pct}% | ≤ 3d (area) = {w3_area_pct}%')
print(f'{"=" * 50}')
print(f'\nGranular Breakdown:')
for label, items, *_ in BUCKET_STYLES:
    print(f'  {label:<25} {len(items):>5} ({pct(len(items))}%)')