import os
import sys
import django

# ── Bootstrap Django — must be before any model imports ──────────────────────
# Change 'allocate.settings' if your settings module has a different name
# e.g. 'worktender.settings' or 'config.settings'
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'allocate.settings')
django.setup()

# ── Imports ───────────────────────────────────────────────────────────────────
from datetime import date
from decimal import Decimal
from tender.models import JobActivity, Allocation

import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side
)
from openpyxl.utils import get_column_letter

# ── Config ────────────────────────────────────────────────────────────────────
CUTOFF_DATE        = date(2026, 3, 13)
OUTPUT_FILE        = 'unallocated_jobs.xlsx'
OUTPUT_COMPLETED   = 'allocations_completed.xlsx'
OUTPUT_PENDING     = 'allocations_not_completed.xlsx'

# ── Query ─────────────────────────────────────────────────────────────────────
print(f"Querying jobs with scheduled_date < {CUTOFF_DATE}, total_area > 0, no allocations...")

print(f"Querying jobs with scheduled_date < {CUTOFF_DATE}, total_area > 0, no allocations...")

allocated_activity_ids = list(Allocation.objects.values_list('job_activity_id', flat=True))

# ── Diagnostics: print counts at each filter step ────────────────────────────
total_acts          = JobActivity.objects.count()
with_date           = JobActivity.objects.filter(scheduled_date__isnull=False).count()
before_cutoff       = JobActivity.objects.filter(scheduled_date__lt=CUTOFF_DATE).count()
with_area           = JobActivity.objects.filter(scheduled_date__lt=CUTOFF_DATE, total_area__gt=0).count()
has_allocs          = JobActivity.objects.filter(id__in=allocated_activity_ids).count()
final               = JobActivity.objects.filter(scheduled_date__lt=CUTOFF_DATE, total_area__gt=0).exclude(id__in=allocated_activity_ids).count()

print(f"\n── Diagnostic Counts ─────────────────────────────────────────")
print(f"  Total JobActivity rows          : {total_acts}")
print(f"  With scheduled_date set         : {with_date}")
print(f"  scheduled_date < {CUTOFF_DATE}  : {before_cutoff}")
print(f"  + total_area > 0                : {with_area}")
print(f"  Total allocated activity IDs    : {len(allocated_activity_ids)}")
print(f"  JobActivities that have allocs  : {has_allocs}")
print(f"  Final (no alloc + area + date)  : {final}")
print(f"──────────────────────────────────────────────────────────────\n")

# ── Also show date range of all JobActivities ─────────────────────────────────
from django.db.models import Min, Max
date_range = JobActivity.objects.filter(
    scheduled_date__isnull=False
).aggregate(min_date=Min('scheduled_date'), max_date=Max('scheduled_date'))
print(f"  Date range in DB: {date_range['min_date']} → {date_range['max_date']}")

# ── Sample: show 5 activities regardless of filters to confirm data ───────────
print(f"\n── Sample 5 JobActivities (no filters) ───────────────────────")
for a in JobActivity.objects.select_related('job','activity','plot')[:5]:
    print(f"  id={a.id} | date={a.scheduled_date} | area={a.total_area} | alloc_status={a.allocation_status} | activity={a.activity.name} | has_allocs={a.id in allocated_activity_ids}")

print()

acts = (
    JobActivity.objects
    .filter(
        scheduled_date__lt=CUTOFF_DATE,
        total_area__gt=0,
    )
    .exclude(id__in=allocated_activity_ids)
    .select_related('job', 'job__farmer', 'activity', 'plot')
    .order_by('job__farmer__farmer_name', 'scheduled_date')
)

rows = []
for act in acts:
    estimated_amount = float(act.total_area or 0) * float(act.rate_per_acre or 0)
    rows.append({
        'Job ID':             act.job.job_id,
        'Farmer Name':        act.job.farmer.farmer_name,
        'Farmer Phone':       act.job.farmer.phone_number or '—',
        'Activity':           act.activity.name,
        'Plot Name':          act.plot.name if act.plot else '—',
        'Plot Code':          act.plot.plot_code if act.plot else '—',
        'Total Area (ac)':    float(act.total_area),
        'Allocated Area (ac)':float(act.allocated_area),
        'Rate/Acre (₹)':      float(act.rate_per_acre),
        'Est. Amount (₹)':    round(estimated_amount, 2),
        'Scheduled Date':     str(act.scheduled_date),
        'Allocation Status':  act.allocation_status,
        'Job Status':         act.job.status,
    })

print(f"Found {len(rows)} records.")

if not rows:
    print("No records found. Exiting.")
    sys.exit(0)

# ── Build Excel ───────────────────────────────────────────────────────────────
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Unallocated Jobs'

HEADERS = list(rows[0].keys())

# Styles
HEADER_FONT    = Font(name='Arial', bold=True, color='FFFFFF', size=11)
HEADER_FILL    = PatternFill('solid', start_color='1A1A2E')   # dark navy
ALT_FILL       = PatternFill('solid', start_color='F0F2F5')   # light grey
WHITE_FILL     = PatternFill('solid', start_color='FFFFFF')
CENTER         = Alignment(horizontal='center', vertical='center', wrap_text=False)
LEFT           = Alignment(horizontal='left',   vertical='center')
MONEY_FMT      = '#,##0.00'
AREA_FMT       = '0.00'
DATE_FMT       = 'DD-MMM-YYYY'

THIN = Side(style='thin', color='D0D0D0')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ── Header row ────────────────────────────────────────────────────────────────
ws.row_dimensions[1].height = 32
for col_idx, header in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font      = HEADER_FONT
    cell.fill      = HEADER_FILL
    cell.alignment = CENTER
    cell.border    = BORDER

# ── Data rows ─────────────────────────────────────────────────────────────────
for row_idx, row in enumerate(rows, start=2):
    fill = ALT_FILL if row_idx % 2 == 0 else WHITE_FILL
    ws.row_dimensions[row_idx].height = 20

    for col_idx, (key, val) in enumerate(row.items(), start=1):
        cell        = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill   = fill
        cell.border = BORDER
        cell.font   = Font(name='Arial', size=10)

        # Column-specific formatting
        if key in ('Total Area (ac)', 'Allocated Area (ac)'):
            cell.number_format = AREA_FMT
            cell.alignment     = CENTER
        elif key in ('Rate/Acre (₹)', 'Est. Amount (₹)'):
            cell.number_format = MONEY_FMT
            cell.alignment     = CENTER
        elif key == 'Scheduled Date':
            cell.alignment = CENTER
        elif key in ('Allocation Status', 'Job Status'):
            cell.alignment = CENTER
            # Color-code status
            if val in ('pending',):
                cell.font = Font(name='Arial', size=10, color='E74C3C')  # red
            elif val in ('fully_allocated', 'completed'):
                cell.font = Font(name='Arial', size=10, color='27AE60')  # green
            elif val in ('partially_allocated',):
                cell.font = Font(name='Arial', size=10, color='F39C12')  # orange
        else:
            cell.alignment = LEFT

# ── Column widths ─────────────────────────────────────────────────────────────
COL_WIDTHS = {
    'Job ID':               16,
    'Farmer Name':          24,
    'Farmer Phone':         16,
    'Activity':             28,
    'Plot Name':            20,
    'Plot Code':            14,
    'Total Area (ac)':      16,
    'Allocated Area (ac)':  18,
    'Rate/Acre (₹)':        16,
    'Est. Amount (₹)':      18,
    'Scheduled Date':       16,
    'Allocation Status':    20,
    'Job Status':           14,
}
for col_idx, header in enumerate(HEADERS, start=1):
    ws.column_dimensions[get_column_letter(col_idx)].width = COL_WIDTHS.get(header, 16)

# ── Summary row ───────────────────────────────────────────────────────────────
summary_row = len(rows) + 2
ws.row_dimensions[summary_row].height = 24

SUMMARY_FILL = PatternFill('solid', start_color='E8F0FE')
SUMMARY_FONT = Font(name='Arial', bold=True, size=10, color='1A1A2E')

# Label in col 1
label_cell            = ws.cell(row=summary_row, column=1, value=f'TOTAL  ({len(rows)} records)')
label_cell.font       = SUMMARY_FONT
label_cell.fill       = SUMMARY_FILL
label_cell.alignment  = LEFT
label_cell.border     = BORDER

# Blank cols 2–6
for c in range(2, 7):
    cell        = ws.cell(row=summary_row, column=c, value='')
    cell.fill   = SUMMARY_FILL
    cell.border = BORDER

# Total Area sum — column 7
total_area_col = HEADERS.index('Total Area (ac)') + 1
tc = ws.cell(
    row=summary_row, column=total_area_col,
    value=f'=SUM({get_column_letter(total_area_col)}2:{get_column_letter(total_area_col)}{len(rows)+1})'
)
tc.font = SUMMARY_FONT; tc.fill = SUMMARY_FILL
tc.number_format = AREA_FMT; tc.alignment = CENTER; tc.border = BORDER

# Est. Amount sum — column 10
est_col = HEADERS.index('Est. Amount (₹)') + 1
ec = ws.cell(
    row=summary_row, column=est_col,
    value=f'=SUM({get_column_letter(est_col)}2:{get_column_letter(est_col)}{len(rows)+1})'
)
ec.font = SUMMARY_FONT; ec.fill = SUMMARY_FILL
ec.number_format = MONEY_FMT; ec.alignment = CENTER; ec.border = BORDER

# Fill remaining summary cells blank
for c in range(1, len(HEADERS) + 1):
    if c not in (1, total_area_col, est_col) and c >= 7:
        cell        = ws.cell(row=summary_row, column=c, value='')
        cell.fill   = SUMMARY_FILL
        cell.border = BORDER

# ── Freeze header row + auto filter ──────────────────────────────────────────
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}1'

# ── Save ──────────────────────────────────────────────────────────────────────
wb.save(OUTPUT_FILE)
print(f"\n✅ Excel saved → {OUTPUT_FILE}")
print(f"   {len(rows)} unallocated job activities")
print(f"   Total area : {sum(r['Total Area (ac)'] for r in rows):.2f} ac")
print(f"   Est. amount: ₹{sum(r['Est. Amount (₹)'] for r in rows):,.2f}")


# ══════════════════════════════════════════════════════════════════════════════
# EXPORT 2 — All Allocations split by work_status
# ══════════════════════════════════════════════════════════════════════════════

def build_allocation_rows(qs):
    rows = []
    for a in qs:
        act  = a.job_activity
        job  = act.job
        plot = act.plot
        rows.append({
            'Allocation ID':        a.id,
            'Job ID':               job.job_id,
            'Farmer Name':          job.farmer.farmer_name,
            'Farmer Phone':         job.farmer.phone_number or '—',
            'Activity':             act.activity.name,
            'Plot Name':            plot.name      if plot else '—',
            'Plot Code':            plot.plot_code if plot else '—',
            'Allocated Date':       str(a.allocated_date),
            'Allocated Area (ac)':  float(a.allocated_area or 0),
            'Actual Area Done (ac)':float(a.actual_area_done) if a.actual_area_done is not None else '—',
            'Admin Override (ac)':  float(a.admin_override_area) if a.admin_override_area is not None else '—',
            'Allocated Workers':    a.allocated_workers,
            'Actual Crew Size':     a.actual_crew_size if a.actual_crew_size is not None else '—',
            'Mukkadam':             a.mukkadam.mukkadam_name,
            'Mukkadam Mobile':      a.mukkadam.mobile_numbers or '—',
            'Rate/Acre (₹)':        float(a.mukkadam_rate or 0),
            'Farmer Rate (₹)':      float(a.farmer_rate or 0),
            'Farmer Amount (₹)':    float(a.farmer_amount or 0),
            'Work Status':          a.work_status,
            'Payment Status':       a.payment_status,
            'Report Submitted':     'Yes' if a.report_submitted else 'No',
            'Farmer Agreed':        'Yes' if a.farmer_agreed is True else ('No' if a.farmer_agreed is False else '—'),
            'Allocation Status':    act.allocation_status,
            'Scheduled Date':       str(act.scheduled_date) if act.scheduled_date else '—',
            'Job Status':           job.status,
        })
    return rows


def write_allocation_excel(rows, filepath, title, header_color):
    if not rows:
        print(f"  No records for {filepath}, skipping.")
        return

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = title

    HEADERS = list(rows[0].keys())

    HEADER_FONT  = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    HEADER_FILL  = PatternFill('solid', start_color=header_color)
    ALT_FILL     = PatternFill('solid', start_color='F0F2F5')
    WHITE_FILL   = PatternFill('solid', start_color='FFFFFF')
    SUMMARY_FILL = PatternFill('solid', start_color='E8F0FE')
    SUMMARY_FONT = Font(name='Arial', bold=True, size=10, color='1A1A2E')
    CENTER       = Alignment(horizontal='center', vertical='center')
    LEFT         = Alignment(horizontal='left',   vertical='center')
    THIN         = Side(style='thin', color='D0D0D0')
    BORDER       = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    MONEY_FMT    = '#,##0.00'
    AREA_FMT     = '0.00'

    # Header row
    ws.row_dimensions[1].height = 32
    for ci, h in enumerate(HEADERS, 1):
        c            = ws.cell(row=1, column=ci, value=h)
        c.font       = HEADER_FONT
        c.fill       = HEADER_FILL
        c.alignment  = CENTER
        c.border     = BORDER

    # Data rows
    for ri, row in enumerate(rows, 2):
        fill = ALT_FILL if ri % 2 == 0 else WHITE_FILL
        ws.row_dimensions[ri].height = 20
        for ci, (key, val) in enumerate(row.items(), 1):
            c        = ws.cell(row=ri, column=ci, value=val)
            c.fill   = fill
            c.border = BORDER
            c.font   = Font(name='Arial', size=10)

            if key in ('Allocated Area (ac)', 'Actual Area Done (ac)', 'Admin Override (ac)'):
                if isinstance(val, float):
                    c.number_format = AREA_FMT
                c.alignment = CENTER
            elif key in ('Rate/Acre (₹)', 'Farmer Rate (₹)', 'Farmer Amount (₹)'):
                c.number_format = MONEY_FMT
                c.alignment     = CENTER
            elif key == 'Work Status':
                c.alignment = CENTER
                if val == 'completed':
                    c.font = Font(name='Arial', size=10, color='27AE60')
                elif val == 'in_progress':
                    c.font = Font(name='Arial', size=10, color='F39C12')
                else:
                    c.font = Font(name='Arial', size=10, color='E74C3C')
            elif key in ('Payment Status', 'Allocation Status', 'Job Status',
                         'Report Submitted', 'Farmer Agreed'):
                c.alignment = CENTER
            else:
                c.alignment = LEFT

    # Column widths
    COL_WIDTHS = {
        'Allocation ID':        14, 'Job ID':               14,
        'Farmer Name':          24, 'Farmer Phone':         16,
        'Activity':             28, 'Plot Name':            20,
        'Plot Code':            14, 'Allocated Date':       16,
        'Allocated Area (ac)':  18, 'Actual Area Done (ac)':20,
        'Admin Override (ac)':  18, 'Allocated Workers':    18,
        'Actual Crew Size':     16, 'Mukkadam':             24,
        'Mukkadam Mobile':      18, 'Rate/Acre (₹)':        16,
        'Farmer Rate (₹)':      16, 'Farmer Amount (₹)':    18,
        'Work Status':          16, 'Payment Status':       16,
        'Report Submitted':     18, 'Farmer Agreed':        14,
        'Allocation Status':    20, 'Scheduled Date':       16,
        'Job Status':           14,
    }
    for ci, h in enumerate(HEADERS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = COL_WIDTHS.get(h, 16)

    # Summary row
    sr   = len(rows) + 2
    ws.row_dimensions[sr].height = 24

    label        = ws.cell(row=sr, column=1, value=f'TOTAL  ({len(rows)} records)')
    label.font   = SUMMARY_FONT
    label.fill   = SUMMARY_FILL
    label.alignment = LEFT
    label.border = BORDER

    for c in range(2, len(HEADERS) + 1):
        cell        = ws.cell(row=sr, column=c, value='')
        cell.fill   = SUMMARY_FILL
        cell.border = BORDER

    # Sum allocated area
    alloc_col = HEADERS.index('Allocated Area (ac)') + 1
    sc = ws.cell(
        row=sr, column=alloc_col,
        value=f'=SUM({get_column_letter(alloc_col)}2:{get_column_letter(alloc_col)}{len(rows)+1})'
    )
    sc.font = SUMMARY_FONT; sc.fill = SUMMARY_FILL
    sc.number_format = AREA_FMT; sc.alignment = CENTER; sc.border = BORDER

    # Sum farmer amount
    amt_col = HEADERS.index('Farmer Amount (₹)') + 1
    ac_ = ws.cell(
        row=sr, column=amt_col,
        value=f'=SUM({get_column_letter(amt_col)}2:{get_column_letter(amt_col)}{len(rows)+1})'
    )
    ac_.font = SUMMARY_FONT; ac_.fill = SUMMARY_FILL
    ac_.number_format = MONEY_FMT; ac_.alignment = CENTER; ac_.border = BORDER

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}1'

    wb.save(filepath)
    print(f"  ✅ {filepath}  ({len(rows)} records)")


# ── Fetch both sets ───────────────────────────────────────────────────────────
print("\n── Exporting Allocation files ────────────────────────────────")

base_qs = (
    Allocation.objects
    .select_related(
        'job_activity', 'job_activity__job', 'job_activity__job__farmer',
        'job_activity__activity', 'job_activity__plot', 'mukkadam'
    )
    .order_by('job_activity__job__farmer__farmer_name', 'allocated_date')
)

completed_rows     = build_allocation_rows(base_qs.filter(work_status='completed'))
not_completed_rows = build_allocation_rows(base_qs.exclude(work_status='completed'))

write_allocation_excel(completed_rows,     OUTPUT_COMPLETED, 'Completed',     '27AE60')  # green header
write_allocation_excel(not_completed_rows, OUTPUT_PENDING,   'Not Completed', 'E74C3C')  # red header

print(f"\n  Completed     : {len(completed_rows)} allocations")
print(f"  Not Completed : {len(not_completed_rows)} allocations")
print("\nDone ✅")